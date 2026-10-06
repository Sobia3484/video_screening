#!/usr/bin/env python
"""Step 2: build the 300-video labeling sample and the labeling workbook.

    python scripts/make_labeling_sample.py                                   # blind sheet (no AI suggestions)
    python scripts/make_labeling_sample.py --prelabels data/labeled/ai_prelabels.csv   # AI-assisted sheet

Writes (data/labeled/):
    labeling_sheet.xlsx / labeling_sheet_ai_assisted.xlsx   what you fill in (no split, no query info)
    sample_selection.csv                                      sample_id, video_uid, stratum, split (dev/test), sampling_weight, seed
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from vidscreen.config import load_settings  # noqa: E402
from vidscreen.screening.sampling import draw_sample  # noqa: E402

LABELS = "relevant,irrelevant,potentially_relevant"
SUBTOPICS = "cpr,neonatal,bls,choking,airway_obstruction,chest_compressions,aed,rescue_breathing,combined,none"
FONT = "Arial"
INPUT_COLS = ("label", "subtopic", "notes")
CONF_FILL = {"high": "E2EFDA", "medium": "FCE4D6", "low": "F4B183"}

BLIND_INSTRUCTIONS = [
    ("How to label (about 1 minute per video, 300 videos)", True),
    ("", False),
    ("1. Go to the 'Labeling' sheet. For each row read: title, description, hashtags, tags, creator, duration (open the link only if needed).", False),
    ("2. Fill ONLY the three yellow columns: label, subtopic, notes. Do not edit the other columns. Do not delete rows.", False),
    ("3. label = relevant | irrelevant | potentially_relevant (dropdown). Rules: docs/screening_criteria.md (v1.1).", False),
    ("4. Core test: can a viewer learn what to do for an infant, child or newborn who needs CPR or is choking?", False),
    ("5. If you are unsure after reading the fields, choose potentially_relevant. Do not guess.", False),
    ("6. subtopic only for relevant videos (dropdown); use 'none' otherwise.", False),
    ("7. Do not use transcripts. Judge every language by the same rules. Save often (Ctrl+S).", False),
    ("8. When finished, save as data/labeled/labeled_sample_300.xlsx", False),
]
AI_INSTRUCTIONS = [
    ("AI-assisted labeling with human verification (about 20-30 seconds per video)", True),
    ("", False),
    ("1. Go to the 'Labeling' sheet. For every row read title, description, hashtags, tags, duration (open the link when in doubt).", False),
    ("2. The grey ai_* columns are only a SUGGESTION made from the same text you see. The AI cannot watch the video. Do not edit them.", False),
    ("3. Fill the three yellow columns for EVERY row: label (dropdown), subtopic (dropdown), notes (optional).", False),
    ("   You decide: type the same label as ai_label only after you have checked it, otherwise choose your own.", False),
    ("4. Take most care where ai_confidence is low or medium (orange) and where ai_label is potentially_relevant.", False),
    ("5. Rules: docs/screening_criteria.md (v1.1). Core test: can a viewer learn what to do for an infant, child or newborn who needs CPR or is choking?", False),
    ("6. If unsure after reading the fields: potentially_relevant. Use subtopic only for relevant videos ('none' otherwise).", False),
    ("7. Do not use transcripts. Judge every language by the same rules. Save often (Ctrl+S).", False),
    ("8. When finished, save as data/labeled/labeled_sample_300.xlsx. Agreement between ai_label and your label is computed later and reported.", False),
]
EXAMPLE = [
    ("", False),
    ("Example row (illustration only, not part of your data):", True),
    ("   title: 'Infant CPR Step-by-Step | Red Cross'   hashtags: infantcpr firstaid", False),
    ("   -> label: relevant      subtopic: cpr      notes: shows two-finger compressions on a manikin", False),
]


def _text(value, limit=None):
    value = "" if pd.isna(value) else str(value)
    if limit and len(value) > limit:
        value = value[:limit].rstrip() + " ...[truncated]"
    return value


def build_workbook(sample: pd.DataFrame, path: Path, prelabels: pd.DataFrame = None) -> None:
    ai = prelabels is not None
    if ai:
        sample = sample.merge(prelabels, on="sample_id", how="left", validate="one_to_one")

    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    for r, (text, bold) in enumerate((AI_INSTRUCTIONS if ai else BLIND_INSTRUCTIONS) + EXAMPLE, 1):
        info.cell(row=r, column=1, value=text).font = Font(name=FONT, bold=bold, size=11)
    info.column_dimensions["A"].width = 150

    # (header, width, getter)
    cols = [
        ("sample_id", 10, lambda r: r.sample_id),
        ("platform", 10, lambda r: r.platform),
        ("title", 55, lambda r: _text(r.title, 400)),
        ("description", 70, lambda r: _text(r.description, 700)),
        ("hashtags", 30, lambda r: _text(r.hashtags).replace("|", " ")),
        ("tags", 30, lambda r: _text(r.tags).replace("|", ", ")),
        ("creator_name", 22, lambda r: _text(r.creator_name)),
        ("duration_min", 12, lambda r: round(float(r.duration_seconds) / 60, 1) if pd.notna(r.duration_seconds) else None),
        ("url", 28, lambda r: _text(r.url)),
    ]
    if ai:
        cols += [
            ("ai_label", 20, lambda r: r.ai_label),
            ("ai_confidence", 13, lambda r: r.ai_confidence),
            ("ai_subtopic", 18, lambda r: r.ai_subtopic),
            ("ai_reason", 55, lambda r: r.ai_reason),
        ]
    cols += [("label", 20, lambda r: None), ("subtopic", 20, lambda r: None), ("notes", 40, lambda r: None),
             ("video_uid", 24, lambda r: r.video_uid)]
    index = {name: i for i, (name, _, _) in enumerate(cols, 1)}

    ws = wb.create_sheet("Labeling")
    yellow = PatternFill("solid", start_color="FFF2CC")
    head_fill = PatternFill("solid", start_color="D9E1F2")
    grey = PatternFill("solid", start_color="EDEDED")
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, (name, width, _) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=FONT, bold=True)
        c.fill = yellow if name in INPUT_COLS else (grey if name.startswith("ai_") else head_fill)
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, row in enumerate(sample.itertuples(index=False), 2):
        for i, (name, _, getter) in enumerate(cols, 1):
            v = getter(row)
            c = ws.cell(row=r, column=i, value=v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"  # never let a title be read as a formula
            is_url = name == "url"
            c.font = Font(name=FONT, size=10, color="0563C1" if is_url else "000000", underline="single" if is_url else None)
            c.alignment = wrap
            if name in INPUT_COLS:
                c.fill = yellow
            elif name.startswith("ai_"):
                c.fill = PatternFill("solid", start_color=CONF_FILL.get(row.ai_confidence, "EDEDED")) if name == "ai_confidence" else grey
            if is_url and v:
                c.hyperlink = v
    last = len(sample) + 1
    for name, options in (("label", LABELS), ("subtopic", SUBTOPICS)):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = get_column_letter(index[name])
        dv.add(f"{letter}2:{letter}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--per-platform", type=int, default=150)
    ap.add_argument("--input", help="Default: data/intermediate/master_cleaned.csv")
    ap.add_argument("--out-dir", help="Default: data/labeled")
    ap.add_argument("--prelabels", help="CSV with sample_id, ai_label, ai_confidence, ai_subtopic, ai_reason")
    args = ap.parse_args()

    s = load_settings()
    src = Path(args.input) if args.input else s.intermediate_dir / "master_cleaned.csv"
    out = Path(args.out_dir) if args.out_dir else s.root / "data" / "labeled"
    df = pd.read_csv(src, dtype={"video_id": str}, encoding="utf-8-sig")
    sample = draw_sample(df, per_platform=args.per_platform, seed=args.seed)

    out.mkdir(parents=True, exist_ok=True)
    sel = sample[["sample_id", "video_uid", "platform", "stratum", "split", "sampling_weight"]].copy()
    sel["seed"] = args.seed
    sel.sort_values("sample_id").to_csv(out / "sample_selection.csv", index=False, encoding="utf-8-sig")
    pre = pd.read_csv(args.prelabels, encoding="utf-8-sig") if args.prelabels else None
    name = "labeling_sheet_ai_assisted.xlsx" if pre is not None else "labeling_sheet.xlsx"
    build_workbook(sample, out / name, pre)

    print(f"Sample: {len(sample)} videos from {len(df)}  (seed {args.seed})")
    print(pd.crosstab(sample["split"], sample["platform"], margins=True).to_string())
    print(f"\nWrote {out/name} and {out/'sample_selection.csv'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
